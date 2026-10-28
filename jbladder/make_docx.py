"""
Build an editable Microsoft Word (.docx) version of the paper from paper/jailbreak_ladder.md.

The point of the .docx is a single, human-editable *master* the author can revise in their own
voice; the LaTeX/PDF outputs (make_pdfs.py) and the eventual journal-template version are generated
*from the revised master*, so there is exactly one source of truth for content.

Renders: a title block, a Word table of contents, Heading 1/2 sections (so Word's navigation pane
and TOC work), justified body text, bullet lists, literally numbered lists (Word's auto-numbering
would otherwise run on across the whole document), display blocks, Markdown pipe tables as native
Word tables with a repeating header row, inline **bold** / *italic* / `code`, @@FIG:name@@ tokens as
embedded figures numbered by FIG_ORDER (matching the PDF and the prose), hanging-indent references,
and page numbers in the footer.

    python -m jbladder.make_docx      # -> results/jbladder_paper.docx
"""

import os
import re

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .make_pdfs import FIG_CAPTIONS, FIG_ORDER

ROOT = os.path.dirname(os.path.dirname(__file__))
RESULTS = os.path.join(ROOT, "results")
FIG = os.path.join(RESULTS, "figures")
PAPER_MD = os.path.join(ROOT, "paper", "jailbreak_ladder.md")
OUT = os.path.join(RESULTS, "jbladder_paper.docx")

FIG_NUM = {name: i + 1 for i, name in enumerate(FIG_ORDER)}
BODY_FONT = "Times New Roman"
CODE_FONT = "Consolas"
TEXT_WIDTH_IN = 6.27          # A4 width minus two 1-inch margins
INK = RGBColor(0x14, 0x16, 0x1A)

_INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+?`|\*[^*\s][^*]*?\*)")
_LINK = re.compile(r"\[([^\]]+)\]\([^)]+\)")
_NUMBERED = re.compile(r"^(\d+)\.\s")
# sections that start on a new page
_PAGE_BREAK_BEFORE = ("References", "Appendix A")


def add_runs(paragraph, text, size=None):
    """Add text to a paragraph, honouring **bold**, *italic*, `code`, and [links](url)->text."""
    text = _LINK.sub(r"\1", text)
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            r = paragraph.add_run(part[2:-2]); r.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = paragraph.add_run(part[1:-1]); r.font.name = CODE_FONT
            r.font.size = Pt((size or 11) - 1.5)
            continue
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r = paragraph.add_run(part[1:-1]); r.italic = True
        else:
            r = paragraph.add_run(part)
        if size:
            r.font.size = Pt(size)


def _field(paragraph, instr, placeholder=""):
    """Insert a Word field (PAGE, TOC, ...) that Word computes on open / F9."""
    run = paragraph.add_run()
    for tag, attrs, text in (("w:fldChar", {"w:fldCharType": "begin"}, None),
                             ("w:instrText", {"xml:space": "preserve"}, instr),
                             ("w:fldChar", {"w:fldCharType": "separate"}, None),
                             ("w:t", {}, placeholder),
                             ("w:fldChar", {"w:fldCharType": "end"}, None)):
        el = OxmlElement(tag)
        for k, v in attrs.items():
            el.set(qn(k), v)
        if text is not None:
            el.text = text
        run._r.append(el)


def _shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def _repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def _is_special(line):
    s = line.strip()
    return (not s or s.startswith("#") or s.startswith("|") or s.startswith(">")
            or s.startswith("@@FIG:") or s == "---" or s.startswith("- ")
            or _NUMBERED.match(s) or s.startswith("```"))


def _emit_table(doc, rows):
    # rows: list of raw "| a | b |" lines; rows[1] is the |---| separator
    def cells(line):
        return [c.strip() for c in line.strip().strip("|").split("|")]
    header = cells(rows[0])
    body = [cells(r) for r in rows[2:]]
    ncol = len(header)
    t = doc.add_table(rows=1, cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    size = 8.5 if ncol >= 6 else 9.5
    hdr = t.rows[0]
    _repeat_header(hdr)
    for j, h in enumerate(header):
        p = hdr.cells[j].paragraphs[0]
        add_runs(p, h, size=size)
        for run in p.runs:
            run.bold = True
        _shade(hdr.cells[j], "E8EBEF")
    for brow in body:
        rc = t.add_row().cells
        for j in range(ncol):
            p = rc[j].paragraphs[0]
            add_runs(p, brow[j] if j < len(brow) else "", size=size)
    # column widths proportional to the longest entry (and never narrower than the longest header
    # word), plus an allowance for cell padding, so short columns don't break words mid-way
    plain = lambda v: re.sub(r"[*`]", "", v)
    weight = []
    for j in range(ncol):
        longest_cell = max(len(plain(r[j])) if j < len(r) else 0 for r in body) if body else 0
        longest_word = max(len(w) for w in plain(header[j]).split() or [""])
        weight.append(min(42, max(longest_cell, longest_word)) + 3)
    t.autofit = False
    for row in t.rows:
        for j, c in enumerate(row.cells):
            c.width = Inches(TEXT_WIDTH_IN * weight[j] / sum(weight))
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def _emit_figure(doc, name):
    path = os.path.join(FIG, name)
    num = FIG_NUM.get(name, 0)
    if os.path.exists(path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(path, width=Inches(6.0))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
    cap.paragraph_format.space_after = Pt(12)
    r = cap.add_run(f"Figure {num}. ")
    r.bold = True
    r.font.size = Pt(9.5)
    r2 = cap.add_run(FIG_CAPTIONS.get(name, ""))
    r2.font.size = Pt(9.5)


def _setup(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name, size, before in (("Title", 18, 0), ("Heading 1", 14, 18), ("Heading 2", 12, 12)):
        st = styles[name]
        st.font.name = BODY_FONT
        fonts = st.element.rPr.rFonts
        fonts.set(qn("w:eastAsia"), BODY_FONT)
        # built-in heading styles point at the theme font, which overrides the name set above
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            fonts.attrib.pop(qn(attr), None)
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = INK
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_with_next = True
    styles["Title"].font.bold = True
    for s in doc.sections:
        s.page_width, s.page_height = Inches(8.27), Inches(11.69)      # A4
        s.left_margin = s.right_margin = Inches(1.0)
        s.top_margin = s.bottom_margin = Inches(1.0)
        fp = s.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _field(fp, "PAGE", "1")
    props = doc.core_properties
    props.title = "The Jailbreak Robustness Ladder"
    props.author = "Arian Nuhan"
    props.keywords = "AI safety; jailbreak; cross-lingual; LLM evaluation"


def build():
    md = open(PAPER_MD, encoding="utf-8").read()
    lines = md.split("\n")
    doc = docx.Document()
    _setup(doc)

    in_refs = False          # inside the References section -> hanging-indent entries
    in_front = True          # before the Abstract -> centred title block
    toc_done = False
    i, n = 0, len(lines)
    while i < n:
        line = lines[i]
        s = line.strip()

        if not s or s == "---":
            i += 1
            continue

        # tables: a | row followed by a |---| separator
        if s.startswith("|") and i + 1 < n and set(lines[i + 1].strip()) <= set("|:- "):
            block = [line]
            j = i + 1
            while j < n and lines[j].strip().startswith("|"):
                block.append(lines[j]); j += 1
            _emit_table(doc, block)
            i = j
            continue

        if s.startswith("@@FIG:"):
            m = re.match(r"@@FIG:([^@]+)@@", s)
            if m:
                _emit_figure(doc, m.group(1).strip())
            i += 1
            continue

        if s.startswith("# "):
            doc.add_paragraph(s[2:].strip(), style="Title")
            i += 1
            continue
        if s.startswith("## ") or s.startswith("### "):
            level = 1 if s.startswith("## ") else 2
            text = s[3:].strip() if level == 1 else s[4:].strip()
            if level == 1:
                in_front = False
                in_refs = text == "References"
                # TOC goes right after the keywords, i.e. before the first numbered section
                if not toc_done and text.startswith("1."):
                    # a plain bold line, not a heading, so the TOC does not list itself
                    title = doc.add_paragraph().add_run("Contents")
                    title.bold = True
                    title.font.size = Pt(14)
                    _field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u',
                           "Right-click here and choose Update Field to build the table of contents.")
                    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                    toc_done = True
                elif text.startswith(_PAGE_BREAK_BEFORE):
                    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            doc.add_heading(text, level=level)
            i += 1
            continue

        if s.startswith(">"):
            # display block: one centred line per non-empty quoted line (definitions, bounds)
            while i < n and lines[i].strip().startswith(">"):
                q = lines[i].strip()[1:].strip()
                if q:
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.space_after = Pt(3)
                    add_runs(p, q)
                i += 1
            continue

        if s.startswith("- ") or _NUMBERED.match(s):
            m = _NUMBERED.match(s)
            text = s[m.end():] if m else s[2:]
            j = i + 1
            # gather wrapped continuation lines (indented)
            while j < n and lines[j].startswith("  ") and lines[j].strip():
                text += " " + lines[j].strip(); j += 1
            if m:
                # literal number + hanging indent: Word's List Number would continue counting
                # across every list in the document
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.35)
                p.paragraph_format.first_line_indent = Inches(-0.25)
                add_runs(p, f"{m.group(1)}.\t{text}")
            else:
                p = doc.add_paragraph(style="List Bullet")
                add_runs(p, text)
            p.paragraph_format.space_after = Pt(3)
            i = j
            continue

        # default: paragraph (accumulate wrapped lines until blank/special)
        block = [s]
        j = i + 1
        while j < n and not _is_special(lines[j]):
            block.append(lines[j].strip()); j += 1
        text = " ".join(block)
        p = doc.add_paragraph()
        if in_refs:
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.first_line_indent = Inches(-0.3)
            p.paragraph_format.space_after = Pt(4)
            add_runs(p, text, size=10)
        elif in_front:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_runs(p, text)
        elif text.startswith("**Table"):
            # table caption: kept on the same page as its table
            p.paragraph_format.keep_with_next = True
            add_runs(p, text, size=10)
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_runs(p, text)
        i = j

    os.makedirs(RESULTS, exist_ok=True)
    doc.save(OUT)
    print(f"wrote {OUT}")
    return OUT


if __name__ == "__main__":
    build()
